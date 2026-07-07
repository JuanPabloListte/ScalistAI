"""Importa comparativas de precios de Construcciones Silicio (Excel multi-corralón).

Formato (idéntico al acopio 2023 ya seedeado): cada hoja es una COMPARATIVA DE
PRECIOS de una obra/rubro, con fecha propia, y N corralones cotizando el mismo
ítem (columnas Unitario/Total por proveedor). Fuente REAL, fechada.

Dos salidas:
  1. CORPUS COMPLETO → backend/data/silicio_corpus_<rango>.json: TODOS los
     puntos (material_raw, unidad, fecha, obra, ubicación, proveedor, precio).
     Nada se pierde; queda auditable en git para curación/ML futura.
  2. INGESTA A material_price_history: SOLO los materiales del CANONICAL_MAP
     (match exacto por nombre normalizado, alta confianza). Cada proveedor =
     un PricePoint dated. El resto NO se ingesta (evita polución del catálogo
     con 230 nombres one-off; requiere curación humana — cantidades embebidas,
     accesorios de plomería/electricidad, unidades ambiguas).

HONESTO: precios tal cual del Excel; nada inventado; idempotente por source.

Uso:
    python -m scripts.import_silicio_excel --dry-run
    python -m scripts.import_silicio_excel
"""
from __future__ import annotations

import argparse
import glob
import json
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

import openpyxl

from app.core.database import SessionLocal
from app.cost_intelligence.application.use_cases.record_material_price import RecordMaterialPrice
from app.cost_intelligence.infrastructure.persistence.price_repository import SqlPriceHistoryRepository
from app.models.material import Material

ORG_ID = 1
IMPORTS_DIR = Path("data/imports")
CORPUS_OUT = Path("data/silicio_corpus_2024-2025.json")

# Filas que NO son materiales (resumen/impuestos de la planilla).
_FOOTER = {
    "SUBTOTAL-1", "SUBTOTAL-2", "DESCUENTO", "COTIZACION DÓLAR/PESO",
    "COND. IVA", "OTROS IMPUESTOS", "TOTALES FINAL", "OBSERVACIONES:",
    "DESIGNACIÓN", "DESIGNACION", "MEJOR PRECIO", "FLETE",
}

# Match EXACTO por nombre normalizado → material canónico existente. Conservador
# a propósito: solo estructurales inequívocos. `units` restringe (ej. arena solo
# si viene por m3, no una compra por "unidad"); `max_price` descarta el caso
# "palet" colado como unitario.
CANONICAL_MAP: dict[str, dict] = {
    "CEMENTO": {"id": 6, "units": {"bolsa", "un", "ud"}, "max_price": 30_000},
    "CAL HIDRATADA": {"id": 36, "units": {"un", "ud", "unid"}, "max_price": 20_000},
    "HERCAL": {"id": 37, "units": {"un", "ud", "unid"}, "max_price": 20_000},
    "ARENA GRUESA": {"id": 7, "units": {"m3", "mt3"}, "max_price": 120_000},
    "LADRILLO COMUN": {"id": 33, "units": {"un", "ud", "unid"}, "max_price": 1_000},
    "PORTANTE 12X19X33": {"id": 30, "units": {"un", "ud", "unid"}, "max_price": 3_000},
}

_MONTHS = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9,
    "octubre": 10, "noviembre": 11, "diciembre": 12,
}


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().upper()


def _norm_unit(u) -> str:
    return _norm(u).lower().replace(".", "") if u else "?"


def _file_year(path: str) -> int:
    m = re.search(r"(20\d{2})", Path(path).name)
    return int(m.group(1)) if m else 2024


def _sheet_date(ws, default_year: int) -> date | None:
    """Fecha de la hoja desde el encabezado 'Córdoba, <día> de <mes> [de <año>]'."""
    for row in ws.iter_rows(min_row=1, max_row=8, values_only=True):
        for cell in (row or []):
            if not cell:
                continue
            txt = str(cell)
            low = _norm(txt).lower()
            for mname, mnum in _MONTHS.items():
                if mname in low:
                    dm = re.search(r"(\d{1,2})\s*de\s*" + mname, low) or re.search(r"(\d{1,2})\s*" + mname, low)
                    day = int(dm.group(1)) if dm else 1
                    ym = re.search(r"(20\d{2})", txt)
                    year = int(ym.group(1)) if ym else default_year
                    try:
                        return date(year, mnum, min(day, 28))
                    except ValueError:
                        return date(year, mnum, 1)
    return None


def _find_header(ws):
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        if row and row[0] and "designaci" in str(row[0]).lower():
            return i, row
    return None, None


def _vendor_cols(header_row) -> list[tuple[int, str]]:
    out = []
    for c, val in enumerate(header_row):
        if not val:
            continue
        name = _norm(val)
        if name in ("DESIGNACION", "UNID", "CANT", "MEJOR PRECIO") or "UNITARIO" in name or "TOTAL" in name:
            continue
        out.append((c, str(val).strip()))
    return out


def _obra(ws) -> str:
    for row in ws.iter_rows(min_row=1, max_row=8, values_only=True):
        for cell in (row or []):
            if cell and "obra:" in str(cell).lower():
                return str(cell).split(":", 1)[1].strip()[:40]
    return "?"


def run(dry: bool) -> int:
    corpus: list[dict] = []
    with SessionLocal() as db:
        repo = SqlPriceHistoryRepository(db)
        record = RecordMaterialPrice(repo)
        ingested = 0
        mapped_names: set[str] = set()

        for path in sorted(glob.glob(str(IMPORTS_DIR / "*.xlsx"))):
            year = _file_year(path)
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            for sh in wb.sheetnames:
                ws = wb[sh]
                hr, hrow = _find_header(ws)
                if hr is None:
                    continue
                sdate = _sheet_date(ws, year) or date(year, 1, 1)
                obra = _obra(ws)
                vcols = _vendor_cols(hrow)
                for row in ws.iter_rows(min_row=hr + 2, values_only=True):
                    if not row or not row[0]:
                        continue
                    raw = str(row[0]).strip()
                    up = _norm(raw)
                    if up in _FOOTER or "COMPARATIVA" in up or not up:
                        continue
                    unit = _norm_unit(row[1]) if len(row) > 1 else "?"
                    for c, vendor in vcols:
                        price = row[c] if c < len(row) else None
                        if not isinstance(price, (int, float)) or price <= 0:
                            continue
                        corpus.append({
                            "material_raw": raw, "unit": unit, "date": sdate.isoformat(),
                            "obra": obra, "vendor": vendor, "price": round(float(price), 2),
                            "file": Path(path).name, "sheet": sh,
                        })
                        # ¿mapea a canónico?
                        spec = CANONICAL_MAP.get(up)
                        if spec and unit in spec["units"] and price <= spec["max_price"]:
                            mapped_names.add(up)
                            if not dry:
                                # source acotado a varchar(64); obra/archivo viven en el corpus JSON.
                                src = f"Silicio {vendor} {sdate.isoformat()}"[:64]
                                # idempotencia: mismo material+fecha+source
                                from sqlalchemy import text
                                dup = db.execute(text(
                                    "select 1 from material_price_history where material_id=:i "
                                    "and date=:d and source=:s limit 1"),
                                    {"i": spec["id"], "d": sdate, "s": src}).first()
                                if not dup:
                                    record.execute(spec["id"], round(float(price), 2),
                                                   source=src, observed_on=sdate)
                                    ingested += 1
            wb.close()
        if not dry:
            db.commit()

        CORPUS_OUT.write_text(json.dumps({
            "source": "Construcciones Silicio — comparativas de precios (Excel)",
            "files": sorted({c["file"] for c in corpus}),
            "date_range": [min(c["date"] for c in corpus), max(c["date"] for c in corpus)] if corpus else [],
            "vendors": sorted({c["vendor"] for c in corpus}),
            "n_points": len(corpus), "points": corpus,
        }, ensure_ascii=False, indent=1), encoding="utf-8")

        dates = sorted({c["date"] for c in corpus})
        print(f"CORPUS: {len(corpus)} puntos de precio, {len({c['material_raw'] for c in corpus})} "
              f"materiales, {len({c['vendor'] for c in corpus})} corralones, "
              f"fechas {dates[0]}…{dates[-1]} → {CORPUS_OUT}")
        print(f"{'[DRY-RUN] ' if dry else ''}INGESTADO a serie histórica: {ingested} puntos "
              f"de {len(mapped_names)} materiales canónicos {sorted(mapped_names)}")
        print("Resto: en el corpus JSON, pendiente de curación humana (no se poluye el catálogo).")
        return 0


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args()
    return run(args.dry_run)


if __name__ == "__main__":
    sys.exit(main())
