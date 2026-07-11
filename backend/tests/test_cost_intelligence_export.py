"""Tests de Fase 5: el `build_workbook` genera un XLSX válido de 4 hojas.
Reabre el archivo con openpyxl y verifica estructura + valores. Sin DB.

Corre directo:  python tests/test_cost_intelligence_export.py
"""
import sys
from io import BytesIO
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parent.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from openpyxl import load_workbook  # noqa: E402

from app.cost_intelligence.infrastructure.export.xlsx_exporter import build_workbook  # noqa: E402

SIM = {
    "name": "Presupuesto Test",
    "totals": {"materials": 1000.0, "labor_cost": 500.0, "labor_hours": 20.0, "duration_days": 2.0},
    "lines": [{"material_name": "Ladrillo", "quantity": 160.0, "unit": "un", "unit_cost": 6.25, "total": 1000.0}],
}
SCEN = [{"entity": "Muro", "base": "Ladrillo", "alternative": "Durlock",
         "material_diff": -200.0, "labor_diff": -100.0, "time_saved": 1.5}]
PROJ = [{"horizon_months": 6, "projected_cost": 1800.0, "variation": 20.0}]


def _wb(scen=SCEN, proj=PROJ):
    return load_workbook(BytesIO(build_workbook(SIM, scen, proj)))


def test_produces_valid_xlsx_with_4_sheets():
    data = build_workbook(SIM, SCEN, PROJ)
    assert isinstance(data, bytes) and len(data) > 0
    assert _wb().sheetnames == ["Presupuesto", "Escenarios", "Proyección", "Detalle materiales"]


def test_presupuesto_costo_directo():
    ws = _wb()["Presupuesto"]
    vals = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
    assert vals.get("Materiales") == 1000.0
    assert vals.get("COSTO DIRECTO") == 1500.0  # materiales + MO


def test_presupuesto_precio_venta_con_breakdown():
    bd = {"overhead": 225.0, "profit": 172.5, "net": 1897.5, "iva": 398.48, "total": 2295.98}
    from io import BytesIO
    from app.cost_intelligence.infrastructure.export.xlsx_exporter import build_workbook
    ws = load_workbook(BytesIO(build_workbook(SIM, SCEN, PROJ, bd)))["Presupuesto"]
    vals = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
    assert vals.get("PRECIO DE VENTA") == 2295.98
    assert vals.get("IVA") == 398.48


def test_detalle_lista_material():
    ws = _wb()["Detalle materiales"]
    assert any(ws.cell(r, 1).value == "Ladrillo" for r in range(1, ws.max_row + 1))


def test_escenarios_muestra_diff():
    ws = _wb()["Escenarios"]
    found = False
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 3).value == "Durlock":
            assert ws.cell(r, 4).value == -200.0
            found = True
    assert found


def test_hojas_vacias_no_rompen():
    wb = _wb(scen=[], proj=[])
    assert "Escenarios" in wb.sheetnames and "Proyección" in wb.sheetnames


# --- Presupuesto completo (budget-summary → Excel) ---

from app.cost_intelligence.infrastructure.export.xlsx_exporter import build_budget_workbook  # noqa: E402

BUDGET = {
    "project_name": "Obra Test",
    "area_m2": 100.0,
    "area_estimated": False,
    "materials_total": 1000.0,
    "labor_total": 500.0,
    "obra_gris_direct": 1500.0,
    "parametric_total": 300.0,
    "direct_cost": 1800.0,
    "sale_price": 2500.0,
    "labor_hours": 40.0,
    "duration_days": 10.0,
    "cost_per_m2": 18.0,
    "breakdown": {
        "direct": 1800.0, "overhead": 270.0, "profit": 200.0, "iva": 230.0, "total": 2500.0,
    },
    "categories": [
        {"name": "Mampostería", "total": 1000.0, "pct": 55.5, "per_m2": 10.0, "parametric": False},
        {"name": "Inst. eléctrica", "total": 300.0, "pct": 16.7, "per_m2": 3.0, "parametric": True},
    ],
    "takeoff": {
        "wall_ml": 40.0, "wall_m2": 112.0, "openings": 5, "doors": 2, "windows": 3,
        "rooms": 4, "floor_m2": 100.0, "roof_m2": 100.0, "columns": 8, "beams_ml": 20.0,
        "cloaca_ml": 0, "electricidad_ml": 0, "sanitarios": 3, "bocas_electricas": 12,
        "escaleras": 0, "pilotes": 0, "zapatas_ml": 0, "armaduras": 0, "armadura_kg": 0,
    },
}


def test_budget_workbook_sheets_and_totals():
    data = build_budget_workbook(BUDGET)
    assert isinstance(data, bytes) and len(data) > 0
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Presupuesto", "Rubros", "Cómputo"]
    ws = wb["Presupuesto"]
    vals = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
    assert vals.get("COSTO DIRECTO") == 1800.0
    assert vals.get("PRECIO DE VENTA") == 2500.0
    assert vals.get("Rubros estimados (paramétricos)") == 300.0


def test_budget_workbook_rubros_marca_estimado():
    wb = load_workbook(BytesIO(build_budget_workbook(BUDGET)))
    ws = wb["Rubros"]
    types = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(4, ws.max_row + 1)}
    assert types.get("Mampostería") == "Modelado"
    assert types.get("Inst. eléctrica") == "Estimado"


def _run():
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    passed = 0
    for t in tests:
        try:
            t()
            print(f"  ok    {t.__name__}")
            passed += 1
        except Exception as e:  # noqa: BLE001
            print(f"  FAIL  {t.__name__}: {e}")
    print(f"\n{passed}/{len(tests)} tests pasaron")
    return passed == len(tests)


if __name__ == "__main__":
    sys.exit(0 if _run() else 1)
