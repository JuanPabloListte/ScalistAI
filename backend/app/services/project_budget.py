"""Presupuesto del proyecto: lista del proveedor guardada × reglas × IFC.

La vista se recalcula siempre en vivo (si cambia el juego de reglas o se
re-procesa el IFC, el presupuesto se actualiza solo); lo único persistido
son el archivo, su lectura y las decisiones del usuario en la revisión.
"""
from __future__ import annotations

import copy
import uuid
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.core.config import settings
from app.models.plan import Plan
from app.models.project import Project
from app.models.project_budget_settings import ProjectBudgetSettings
from app.models.project_price_list import ProjectPriceList
from app.models.rule_set import RuleSet
from app.services import budget_rules as br
from app.services.price_list_parse import parse_price_list


def model_status(project_id: int, db: Session) -> str:
    """Estado del último IFC del proyecto: ready / processing / error / missing."""
    plans = db.scalars(
        select(Plan)
        .where(Plan.project_id == project_id, Plan.original_filename.ilike("%.ifc"))
        .order_by(Plan.id.desc())
    ).all()
    if not plans:
        return "missing"
    if any(p.status == "ready" for p in plans):
        return "ready"
    if plans[0].status == "processing":
        return "processing"
    return "error"


def resolve_rule_set(project: Project, db: Session, rule_set_id: int | None = None) -> RuleSet:
    """El juego pedido, el del proyecto, o el estándar del sistema. Sirven los
    de la org del proyecto y el del sistema."""
    from sqlalchemy import or_

    system = br.ensure_system_rule_set(db)
    for rid in (rule_set_id, project.rule_set_id):
        if rid is None:
            continue
        rs = db.scalar(select(RuleSet).where(
            RuleSet.id == rid,
            or_(RuleSet.organization_id == project.organization_id, RuleSet.is_system.is_(True)))
            .options(selectinload(RuleSet.rules)))
        if rs is not None:
            return rs
        if rid == rule_set_id:
            raise LookupError("Juego de reglas no encontrado")
    return system


def get_price_list(project_id: int, db: Session) -> ProjectPriceList | None:
    return db.scalar(select(ProjectPriceList).where(ProjectPriceList.project_id == project_id))


def _patched_lines(pl: ProjectPriceList) -> list[dict]:
    lines = copy.deepcopy(pl.parsed.get("lines") or [])
    for ln in lines:
        patch = (pl.line_patches or {}).get(str(ln["row"]))
        if patch:
            pres = ln.setdefault("presentation", {"attrs": {}})
            pres["amount"] = float(patch["amount"])
            pres["unit"] = patch["unit"]
            ln["patched"] = True
    return lines


def get_budget_settings(project_id: int, db: Session) -> ProjectBudgetSettings | None:
    return db.scalar(select(ProjectBudgetSettings)
                     .where(ProjectBudgetSettings.project_id == project_id))


def update_budget_settings(project: Project, db: Session, *, labor: dict | None = None,
                           overhead_pct: float | None = None, profit_pct: float | None = None,
                           reset_margins: bool = False) -> None:
    """`labor`: rule_id → {"mode", "unit_price"}; un valor null borra el ajuste
    de ese rubro (vuelve a lo que dice la regla)."""
    s = get_budget_settings(project.id, db)
    if s is None:
        s = ProjectBudgetSettings(project_id=project.id, labor={})
        db.add(s)
    if labor is not None:
        merged = dict(s.labor or {})
        for key, val in labor.items():
            if val is None:
                merged.pop(str(key), None)
            else:
                merged[str(key)] = {k: v for k, v in val.items() if k in ("mode", "unit_price")}
        s.labor = merged
    if reset_margins:
        s.overhead_pct = s.profit_pct = None
    if overhead_pct is not None:
        s.overhead_pct = overhead_pct
    if profit_pct is not None:
        s.profit_pct = profit_pct
    db.flush()


def _org_margins(organization_id: int, db: Session) -> tuple[float, float, float]:
    from app.cost_intelligence.infrastructure.persistence.cost_settings_repo import (
        SqlCostSettingsRepo,
    )

    cs = SqlCostSettingsRepo(db).get_or_create(organization_id)
    return cs.overhead_pct, cs.profit_pct, cs.iva_pct


def build_view(project: Project, db: Session) -> dict:
    from app.services.labor_rates import trade_costs

    rule_set = resolve_rule_set(project, db)
    pl = get_price_list(project.id, db)
    status = model_status(project.id, db)
    settings_row = get_budget_settings(project.id, db)
    org_overhead, org_profit, iva_pct = _org_margins(project.organization_id, db)
    overhead = settings_row.overhead_pct if settings_row and settings_row.overhead_pct is not None \
        else org_overhead
    profit = settings_row.profit_pct if settings_row and settings_row.profit_pct is not None \
        else org_profit
    view: dict = {
        "rule_set": {"id": rule_set.id, "name": rule_set.name} if rule_set else None,
        "model_status": status,
        "price_list": None,
        "measures": None,
        "budget": None,
        "labor": None,
        "price": None,
        "settings": {
            "overhead_pct": overhead, "profit_pct": profit, "iva_pct": iva_pct,
            "margins_custom": bool(settings_row and (settings_row.overhead_pct is not None
                                                     or settings_row.profit_pct is not None)),
            "org_overhead_pct": org_overhead, "org_profit_pct": org_profit,
            "labor": (settings_row.labor if settings_row else None) or {},
        },
    }
    measures = br.project_measures(project.id, db) if status == "ready" else None
    if measures is not None:
        view["measures"] = measures.as_dict()
        if rule_set is not None:
            costs = trade_costs(project.organization_id, db)
            view["labor"] = br.compute_labor(
                measures, [br.RuleSpec.of(r) for r in rule_set.rules], costs,
                view["settings"]["labor"])
            view["trades"] = [{"trade": k, "daily_cost": v} for k, v in costs.items()]
    if pl is not None:
        view["price_list"] = _price_list_view(pl)
        if measures is not None and rule_set is not None:
            view["budget"] = br.compute_budget(
                measures, [br.RuleSpec.of(r) for r in rule_set.rules],
                view["price_list"]["lines"], overrides=pl.overrides or {})
    if view["labor"] is not None or view["budget"] is not None:
        view["price"] = br.compute_price(
            (view["budget"] or {}).get("totals", {}).get("net", 0.0),
            (view["labor"] or {}).get("total", 0.0),
            overhead_pct=overhead, profit_pct=profit, iva_pct=iva_pct)
    return view


def _price_list_view(pl: ProjectPriceList) -> dict:
    lines = _patched_lines(pl)
    return {
        **{k: v for k, v in pl.parsed.items() if k != "lines"},
        "lines": lines,
        "filename": pl.original_filename,
        "iva_mode": "auto" if pl.iva_included is None else (
            "included" if pl.iva_included else "excluded"),
        "overrides": pl.overrides or {},
        "line_patches": pl.line_patches or {},
        "updated_at": pl.updated_at.isoformat() if pl.updated_at else None,
    }


def _store_file(project_id: int, filename: str, data: bytes) -> str:
    folder = Path(settings.STORAGE_DIR) / "price_lists" / str(project_id)
    folder.mkdir(parents=True, exist_ok=True)
    safe = "".join(c if c.isalnum() or c in "._-" else "_" for c in Path(filename).name)[-120:]
    path = folder / f"{uuid.uuid4().hex[:8]}_{safe}"
    path.write_bytes(data)
    return str(path)


def save_price_list(project: Project, filename: str, data: bytes, db: Session, *,
                    rule_set_id: int | None, iva_included: bool | None) -> ProjectPriceList:
    """Sube/reemplaza la lista del proyecto. Reemplazar borra las decisiones
    de la revisión anterior (eran sobre otros renglones)."""
    parsed = parse_price_list(filename, data, iva_included=iva_included)  # ValueError → 400
    if rule_set_id is not None:
        project.rule_set_id = resolve_rule_set(project, db, rule_set_id).id
    elif project.rule_set_id is None:
        rs = resolve_rule_set(project, db)
        project.rule_set_id = rs.id if rs else None

    pl = get_price_list(project.id, db)
    old_path = pl.file_path if pl else None
    if pl is None:
        pl = ProjectPriceList(project_id=project.id)
        db.add(pl)
    pl.original_filename = filename
    pl.file_path = _store_file(project.id, filename, data)
    pl.iva_included = iva_included
    pl.parsed = parsed
    pl.overrides = {}
    pl.line_patches = {}
    project.wizard_step = max(project.wizard_step, 3)
    db.flush()
    if old_path and old_path != pl.file_path:
        Path(old_path).unlink(missing_ok=True)
    return pl


def update_price_list(project: Project, pl: ProjectPriceList, db: Session, *,
                      rule_set_id: int | None = None, iva_mode: str | None = None,
                      overrides: dict | None = None, line_patches: dict | None = None) -> None:
    if rule_set_id is not None:
        project.rule_set_id = resolve_rule_set(project, db, rule_set_id).id
    if iva_mode is not None:
        forced = {"auto": None, "included": True, "excluded": False}[iva_mode]
        if forced != pl.iva_included:
            data = Path(pl.file_path).read_bytes()
            pl.parsed = parse_price_list(pl.original_filename, data, iva_included=forced)
            pl.iva_included = forced
    if overrides is not None:
        pl.overrides = overrides
    if line_patches is not None:
        pl.line_patches = line_patches
    db.flush()


# ---------------------------------------------------------------------------
# Exportación a Excel
# ---------------------------------------------------------------------------

_UNIT_LABEL = {"kg": "kg", "m3": "m³", "m2": "m²", "m": "m", "l": "l", "u": "u"}
_STATUS_LABEL = {"ok": "Con precio", "sin_precio": "Sin precio",
                 "sin_medida": "El modelo no lo trae",
                 "unidad_incompatible": "Unidad incompatible",
                 "regla_invalida": "Regla inválida"}


_LABOR_STATUS = {"ok": "Con precio", "sin_precio": "Falta el precio",
                 "sin_jornal": "Faltan jornales", "sin_rendimiento": "Falta el rendimiento",
                 "sin_medida": "El modelo no lo trae"}


def budget_xlsx(project: Project, view: dict) -> bytes:
    """Presupuesto de obra: resumen con precio de venta, lista de compras, mano
    de obra, detalle de materiales por regla y medidas del modelo. `view` es la
    salida de build_view (con budget)."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    budget, pl, measures = view["budget"], view["price_list"], view["measures"]
    labor, price = view.get("labor"), view.get("price")
    money = '"$" #,##0.00'
    qty = "#,##0.00"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1E293B")
    unit = lambda u: _UNIT_LABEL.get(u, u)  # noqa: E731

    wb = Workbook()

    def table(ws, headers, rows, widths, formats):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append(r)
            for col, fmt in formats.items():
                ws.cell(row=ws.max_row, column=col).number_format = fmt
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    # --- Resumen
    ws = wb.active
    ws.title = "Resumen"
    t = budget["totals"]
    ws.append(["Presupuesto de obra"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    info = [
        ("Proyecto", project.name),
        ("Proveedor", pl.get("supplier") or "—"),
        ("Archivo", pl.get("filename")),
        ("Fecha del presupuesto", pl.get("document_date") or "—"),
        ("Reglas de cómputo", (view.get("rule_set") or {}).get("name", "—")),
        ("Precios del archivo",
         "con IVA (se pasaron a neto)" if pl["iva"]["included"] else "sin IVA"),
        None,
        ("Materiales (neto)", t["net"]),
    ]
    if price is not None:
        info += [
            ("Mano de obra", price["labor"]),
            ("Costo directo", price["direct"]),
            (f"Gastos generales {round(price['overhead_pct'] * 100, 1):g}%", price["overhead"]),
            (f"Beneficio {round(price['profit_pct'] * 100, 1):g}%", price["profit"]),
            ("Precio sin IVA", price["net"]),
            (f"IVA {round(price['iva_pct'] * 100)}%", price["iva"]),
            ("Precio de venta", price["total"]),
        ]
    else:
        info += [(f"IVA {round(t['iva_rate'] * 100)}%", t["iva"]),
                 ("Total con IVA", t["total"])]
    info += [
        None,
        ("Ítems con precio", budget["counts"]["ok"]),
        ("Ítems sin precio", budget["counts"]["sin_precio"]),
        ("Ítems sin medida en el modelo", budget["counts"]["sin_medida"]),
    ]
    for entry in info:
        if entry is None:
            ws.append([])
            continue
        ws.append(list(entry))
        if isinstance(entry[1], float):
            ws.cell(row=ws.max_row, column=2).number_format = money
        if entry[0] in ("Total con IVA", "Precio de venta"):
            for col in (1, 2):
                ws.cell(row=ws.max_row, column=col).font = Font(bold=True)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 48

    # --- Lista de compras
    rows = [[s.get("code") or "", s["description"],
             f"{s['presentation']['amount']:g} {unit(s['presentation']['unit'])}",
             s["units_to_buy"], s["unit_price_net"], s["subtotal"]]
            for s in budget["shopping_list"]]
    rows.append(["", "TOTAL NETO", "", None, None, t["net"]])
    table(wb.create_sheet("Lista de compras"),
          ["Código", "Producto del proveedor", "Presentación", "Unidades",
           "Precio unit. neto", "Subtotal neto"],
          rows, [14, 60, 14, 12, 18, 18], {4: "#,##0", 5: money, 6: money})

    # --- Detalle por regla
    rows = []
    for it in budget["items"]:
        status = _STATUS_LABEL.get(it.get("status"), it.get("status") or "")
        if it.get("manual"):
            status += " (a mano)"
        rows.append([
            it.get("measure_label"), it["label"], it.get("measure_value"),
            unit(it.get("measure_unit")), it.get("factor"), it.get("waste_pct"),
            it.get("quantity"), unit(it.get("unit")),
            (it.get("line") or {}).get("description", ""),
            it.get("units_to_buy"), it.get("subtotal"), status,
        ])
    table(wb.create_sheet("Detalle"),
          ["Medida", "Producto", "Medida del modelo", "Unidad", "Factor", "Desperdicio %",
           "Cantidad", "Unidad", "Renglón del proveedor", "Unidades", "Subtotal neto", "Estado"],
          rows, [30, 34, 14, 8, 10, 12, 14, 8, 50, 12, 16, 18],
          {3: qty, 5: "0.0000", 7: qty, 10: qty, 11: money})

    # --- Mano de obra
    if labor is not None:
        mode_label = {"unit": "Tercerizado", "daily": "Cuadrilla propia"}
        rows = []
        for it in labor["items"]:
            how = mode_label.get(it["mode"], it["mode"])
            if it["mode"] == "daily" and it.get("workdays"):
                how += f" · {it['workdays']:g} jornales"
            rows.append([
                it["stage"], it["label"], it["quantity"], unit(it["measure_unit"]), how,
                it.get("days"), it.get("unit_price"), it.get("cost"),
                _LABOR_STATUS.get(it["status"], it["status"]),
            ])
        rows.append(["", "TOTAL MANO DE OBRA", None, "", "", None, None, labor["total"], ""])
        table(wb.create_sheet("Mano de obra", 2),
              ["Rubro", "Tarea", "Cantidad", "Unidad", "Cómo se cotiza", "Días",
               "Precio por unidad", "Subtotal", "Estado"],
              rows, [16, 40, 12, 8, 30, 10, 18, 18, 20], {3: qty, 6: "0.0", 7: money, 8: money})

    # --- Medidas del modelo
    rows = [[m["label"], m["value"], unit(m["unit"]), m["count"]] for m in measures["measures"]]
    rows += [[f"  Hierro Ø{r['diameter_mm']:g}", r["kg"], "kg", None]
             for r in measures["rebar_by_diameter"]]
    rows += [[f"Aviso: {w}", None, "", None] for w in measures["warnings"]]
    table(wb.create_sheet("Medidas del modelo"), ["Medida", "Valor", "Unidad", "Elementos"],
          rows, [60, 14, 8, 10], {2: qty})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
