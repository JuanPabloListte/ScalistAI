"""Genera el XLSX de 4 hojas del Motor de Inteligencia de Costos (Módulo 5).

Toma estructuras de datos planas (dicts/listas, no objetos de dominio) → es
fácil de testear y desacoplado. Hojas:
  1. Presupuesto      — totales (materiales, MO, duración, total)
  2. Escenarios       — comparativa de recetas alternativas (Δ$ y días)
  3. Proyección       — costo proyectado a futuro (IPC real)
  4. Detalle materiales — cantidades y costos por material
"""
from __future__ import annotations

from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=14)
_TOTAL_FONT = Font(bold=True)
_MONEY_FMT = '#,##0.00 "ARS"'
_PCT_FMT = '+0.0"%";-0.0"%"'


def _header_row(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(sim: dict, scenarios: Iterable[dict], projections: Iterable[dict],
                   breakdown: dict | None = None) -> bytes:
    wb = Workbook()
    totals = sim["totals"]
    total_obra = totals["materials"] + totals["labor_cost"]

    # --- Hoja 1: Presupuesto ---
    ws = wb.active
    ws.title = "Presupuesto"
    ws["A1"] = sim.get("name") or "Presupuesto"
    ws["A1"].font = _TITLE_FONT
    _header_row(ws, 3, ["Rubro", "Monto"])
    # Costo directo + (si hay) markup a precio de venta + métricas.
    rows: list[tuple[str, float, str]] = [
        ("Materiales", totals["materials"], "money"),
        ("Mano de obra", totals["labor_cost"], "money"),
        ("COSTO DIRECTO", total_obra, "bold"),
    ]
    if breakdown:
        rows += [
            ("Gastos generales", breakdown["overhead"], "money"),
            ("Beneficio", breakdown["profit"], "money"),
            ("Precio neto (sin IVA)", breakdown["net"], "money"),
            ("IVA", breakdown["iva"], "money"),
            ("PRECIO DE VENTA", breakdown["total"], "bold"),
        ]
    rows += [
        ("Horas hombre", totals["labor_hours"], "plain"),
        ("Duración estimada (días)", totals["duration_days"], "plain"),
    ]
    for i, (label, value, kind) in enumerate(rows, start=4):
        lc = ws.cell(row=i, column=1, value=label)
        c = ws.cell(row=i, column=2, value=round(value, 2))
        if kind != "plain":
            c.number_format = _MONEY_FMT
        if kind == "bold":
            lc.font = _TOTAL_FONT
            c.font = _TOTAL_FONT
    _autosize(ws, [28, 22])

    # --- Hoja 2: Escenarios ---
    ws2 = wb.create_sheet("Escenarios")
    ws2["A1"] = "Comparativa de sistemas alternativos"
    ws2["A1"].font = _TITLE_FONT
    _header_row(ws2, 3, ["Entidad", "Sistema base", "Alternativa",
                         "Δ Materiales", "Δ Mano de obra", "Días ahorrados"])
    r = 4
    for s in scenarios:
        ws2.cell(row=r, column=1, value=s["entity"])
        ws2.cell(row=r, column=2, value=s["base"])
        ws2.cell(row=r, column=3, value=s["alternative"])
        ws2.cell(row=r, column=4, value=round(s["material_diff"], 2)).number_format = _MONEY_FMT
        ws2.cell(row=r, column=5, value=round(s["labor_diff"], 2)).number_format = _MONEY_FMT
        ws2.cell(row=r, column=6, value=round(s["time_saved"], 1))
        r += 1
    if r == 4:
        ws2.cell(row=4, column=1, value="(sin entidades con alternativas en este plano)")
    _autosize(ws2, [16, 30, 30, 18, 18, 16])

    # --- Hoja 3: Proyección ---
    ws3 = wb.create_sheet("Proyección")
    ws3["A1"] = "Proyección de costo (inflación IPC)"
    ws3["A1"].font = _TITLE_FONT
    _header_row(ws3, 3, ["Horizonte", "Costo proyectado", "Variación"])
    ws3.cell(row=4, column=1, value="Hoy")
    ws3.cell(row=4, column=2, value=round(total_obra, 2)).number_format = _MONEY_FMT
    r = 5
    for p in projections:
        ws3.cell(row=r, column=1, value=f"{p['horizon_months']} meses")
        ws3.cell(row=r, column=2, value=round(p["projected_cost"], 2)).number_format = _MONEY_FMT
        ws3.cell(row=r, column=3, value=round(p["variation"], 1)).number_format = _PCT_FMT
        r += 1
    if r == 5:
        ws3.cell(row=5, column=1, value="(sin índice de inflación cargado)")
    _autosize(ws3, [16, 22, 14])

    # --- Hoja 4: Detalle materiales ---
    ws4 = wb.create_sheet("Detalle materiales")
    ws4["A1"] = "Detalle de materiales"
    ws4["A1"].font = _TITLE_FONT
    _header_row(ws4, 3, ["Material", "Cantidad", "Unidad", "Precio unit.", "Total"])
    r = 4
    for ln in sim["lines"]:
        ws4.cell(row=r, column=1, value=ln["material_name"])
        ws4.cell(row=r, column=2, value=round(ln["quantity"], 2))
        ws4.cell(row=r, column=3, value=ln["unit"])
        ws4.cell(row=r, column=4, value=round(ln["unit_cost"], 2)).number_format = _MONEY_FMT
        ws4.cell(row=r, column=5, value=round(ln["total"], 2)).number_format = _MONEY_FMT
        r += 1
    _autosize(ws4, [34, 12, 10, 16, 18])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_budget_workbook(summary: dict) -> bytes:
    """XLSX alineado a `ProjectBudgetSummary` (pantalla de presupuesto).

    Hojas:
      1. Presupuesto — obra gris + paramétricos + precio de venta
      2. Rubros — desglose por categoría (modelado vs estimado)
      3. Cómputo — takeoff de cantidades
    """
    wb = Workbook()
    bd = summary.get("breakdown") or {}
    takeoff = summary.get("takeoff") or {}
    categories = summary.get("categories") or []

    # --- Hoja 1: Presupuesto ---
    ws = wb.active
    ws.title = "Presupuesto"
    ws["A1"] = summary.get("project_name") or "Presupuesto"
    ws["A1"].font = _TITLE_FONT
    if summary.get("area_estimated"):
        ws["A2"] = f"Área cubierta ≈ {round(summary.get('area_m2') or 0, 1)} m² (estimada)"
    else:
        ws["A2"] = f"Área cubierta: {round(summary.get('area_m2') or 0, 1)} m²"

    _header_row(ws, 4, ["Rubro", "Monto"])
    rows: list[tuple[str, float, str]] = [
        ("Materiales (obra gris)", float(summary.get("materials_total") or 0), "money"),
        ("Mano de obra (obra gris)", float(summary.get("labor_total") or 0), "money"),
        ("Obra gris (modelado)", float(summary.get("obra_gris_direct") or 0), "bold"),
        ("Rubros estimados (paramétricos)", float(summary.get("parametric_total") or 0), "money"),
        ("COSTO DIRECTO", float(summary.get("direct_cost") or 0), "bold"),
        ("Gastos generales", float(bd.get("overhead") or 0), "money"),
        ("Beneficio", float(bd.get("profit") or 0), "money"),
        ("IVA", float(bd.get("iva") or 0), "money"),
        ("PRECIO DE VENTA", float(summary.get("sale_price") or bd.get("total") or 0), "bold"),
        ("Horas hombre", float(summary.get("labor_hours") or 0), "plain"),
        ("Duración estimada (días)", float(summary.get("duration_days") or 0), "plain"),
    ]
    if summary.get("cost_per_m2") is not None:
        rows.append(("Costo directo / m²", float(summary["cost_per_m2"]), "money"))

    for i, (label, value, kind) in enumerate(rows, start=5):
        lc = ws.cell(row=i, column=1, value=label)
        c = ws.cell(row=i, column=2, value=round(value, 2))
        if kind != "plain":
            c.number_format = _MONEY_FMT
        if kind == "bold":
            lc.font = _TOTAL_FONT
            c.font = _TOTAL_FONT
    _autosize(ws, [36, 22])

    # --- Hoja 2: Rubros ---
    ws2 = wb.create_sheet("Rubros")
    ws2["A1"] = "Desglose por rubro"
    ws2["A1"].font = _TITLE_FONT
    _header_row(ws2, 3, ["Rubro", "Tipo", "Monto", "%", "$/m²"])
    r = 4
    for cat in categories:
        ws2.cell(row=r, column=1, value=cat.get("name"))
        ws2.cell(row=r, column=2, value="Estimado" if cat.get("parametric") else "Modelado")
        ws2.cell(row=r, column=3, value=round(float(cat.get("total") or 0), 2)).number_format = _MONEY_FMT
        ws2.cell(row=r, column=4, value=round(float(cat.get("pct") or 0), 1))
        per_m2 = cat.get("per_m2")
        if per_m2 is not None:
            ws2.cell(row=r, column=5, value=round(float(per_m2), 2)).number_format = _MONEY_FMT
        r += 1
    if r == 4:
        ws2.cell(row=4, column=1, value="(sin rubros)")
    _autosize(ws2, [32, 12, 18, 8, 14])

    # --- Hoja 3: Cómputo ---
    ws3 = wb.create_sheet("Cómputo")
    ws3["A1"] = "Resumen de cómputo (takeoff)"
    ws3["A1"].font = _TITLE_FONT
    _header_row(ws3, 3, ["Cantidad", "Valor", "Detalle"])
    takeoff_rows = [
        ("Muros", f"{round(float(takeoff.get('wall_ml') or 0), 1)} ml",
         f"{round(float(takeoff.get('wall_m2') or 0), 1)} m²"),
        ("Aberturas", str(takeoff.get("openings") or 0),
         f"{takeoff.get('doors') or 0} puertas · {takeoff.get('windows') or 0} ventanas"),
        ("Ambientes", str(takeoff.get("rooms") or 0),
         f"{round(float(takeoff.get('floor_m2') or 0), 1)} m² de piso"),
        ("Cubierta / losa", f"{round(float(takeoff.get('roof_m2') or 0), 1)} m²", ""),
        ("Columnas", str(takeoff.get("columns") or 0), ""),
        ("Vigas", f"{round(float(takeoff.get('beams_ml') or 0), 1)} ml",
         f"+ {takeoff.get('escaleras') or 0} escaleras" if takeoff.get("escaleras") else ""),
        ("Sanitarios", str(takeoff.get("sanitarios") or 0),
         f"{round(float(takeoff.get('cloaca_ml') or 0), 1)} ml cañería"),
        ("Eléctrico", str(takeoff.get("bocas_electricas") or 0),
         f"{round(float(takeoff.get('electricidad_ml') or 0), 1)} ml tendido"),
        ("Pilotes", str(takeoff.get("pilotes") or 0), ""),
        ("Zapatas", f"{round(float(takeoff.get('zapatas_ml') or 0), 1)} ml", ""),
        ("Armaduras", str(takeoff.get("armaduras") or 0),
         f"{round(float(takeoff.get('armadura_kg') or 0), 1)} kg"),
    ]
    r = 4
    for label, value, detail in takeoff_rows:
        ws3.cell(row=r, column=1, value=label)
        ws3.cell(row=r, column=2, value=value)
        ws3.cell(row=r, column=3, value=detail)
        r += 1
    _autosize(ws3, [18, 18, 36])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
