import math
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.deps import get_current_user
from app.models import DetectedElement, Plan, User
from app.schemas.material import MaterialSummaryItem

router = APIRouter(tags=["plans"])

# Subtipos de abertura que restan al PERÍMETRO del muro (zócalo, etc).
# Las ventanas solo restan al área.
_OPENING_SUBTYPES_SUBTRACT_PERIMETER = {"door", "sliding-door"}


def _opening_overlaps_wall(
    opening: DetectedElement, wall: DetectedElement, scale_px_per_m: float
) -> bool:
    """Devuelve True si el centro de la abertura proyecta dentro del segmento de muro."""
    o_pts = opening.geometry.get("points") if opening.geometry else None
    w_pts = wall.geometry.get("points") if wall.geometry else None
    if not o_pts or len(o_pts) < 4 or not w_pts or len(w_pts) < 4:
        return False
    ox = (float(o_pts[0]) + float(o_pts[2])) / 2.0
    oy = (float(o_pts[1]) + float(o_pts[3])) / 2.0
    x1, y1, x2, y2 = float(w_pts[0]), float(w_pts[1]), float(w_pts[2]), float(w_pts[3])
    dx, dy = x2 - x1, y2 - y1
    l2 = dx * dx + dy * dy
    if l2 <= 0:
        return False
    t = ((ox - x1) * dx + (oy - y1) * dy) / l2
    if t < 0.0 or t > 1.0:
        return False
    proj_x = x1 + t * dx
    proj_y = y1 + t * dy
    dist_px = math.hypot(ox - proj_x, oy - proj_y)
    return dist_px <= 0.5 * scale_px_per_m


def _get_materials_summary_data(plan_id: int, page: int | None, db: Session) -> list[dict]:
    plan = db.get(Plan, plan_id)
    page_scales: dict[str, float] = (plan.page_scales or {}) if plan else {}

    stmt = select(DetectedElement).where(DetectedElement.plan_id == plan_id)
    if page is not None:
        stmt = stmt.where(DetectedElement.page == page)
    elements = list(db.scalars(stmt).all())

    openings_by_page: dict[int, list[DetectedElement]] = {}
    for el in elements:
        if el.type == "opening":
            openings_by_page.setdefault(el.page, []).append(el)

    summary_dict = {}

    for element in elements:
        for assembly in element.assemblies:
            for am in assembly.assembly_materials:
                material = am.material
                qty = 0.0
                if assembly.applies_to == "wall" and element.type == "wall":
                    length = element.length_m or 0.0
                    height = element.height_m or 2.8
                    area = length * height
                    scale_px_per_m = page_scales.get(str(element.page))
                    if scale_px_per_m and scale_px_per_m > 0:
                        for op in openings_by_page.get(element.page, []):
                            if _opening_overlaps_wall(op, element, scale_px_per_m):
                                area -= (op.length_m or 0.0) * (op.height_m or 2.1)
                        area = max(area, 0.0)
                    qty = area * am.consumption * (1.0 + am.waste_factor)
                elif element.type == "room":
                    if assembly.applies_to == "room_floor":
                        qty = (element.area_m2 or 0.0) * am.consumption * (1.0 + am.waste_factor)
                    elif assembly.applies_to == "room_wall":
                        area = (element.length_m or 0.0) * (element.height_m or 2.8)
                        qty = area * am.consumption * (1.0 + am.waste_factor)
                    elif assembly.applies_to == "room_perimeter":
                        qty = (element.length_m or 0.0) * am.consumption * (1.0 + am.waste_factor)
                elif element.type == "opening":
                    if assembly.applies_to == "opening":
                        area = (element.length_m or 0.0) * (element.height_m or 2.1)
                        qty = area * am.consumption * (1.0 + am.waste_factor)
                    elif assembly.applies_to == "opening_perimeter":
                        qty = (element.length_m or 0.0) * am.consumption * (1.0 + am.waste_factor)
                elif element.type == "beam" and assembly.applies_to == "beam":
                    qty = (element.length_m or 0.0) * am.consumption * (1.0 + am.waste_factor)
                elif element.type == "roof" and assembly.applies_to == "roof":
                    qty = (element.area_m2 or 0.0) * am.consumption * (1.0 + am.waste_factor)
                elif element.type == "column" and assembly.applies_to == "column":
                    qty = (element.area_m2 or 0.0) * am.consumption * (1.0 + am.waste_factor)

                if qty > 0.0:
                    if material.id not in summary_dict:
                        summary_dict[material.id] = {
                            "material": material,
                            "quantity": 0.0,
                            "unit": material.unit,
                            "unit_price": material.unit_price,
                        }
                    summary_dict[material.id]["quantity"] += qty

    result = []
    for item in summary_dict.values():
        qty = round(item["quantity"], 2)
        unit_price = item.get("unit_price", 0.0)
        result.append({
            "material": item["material"],
            "quantity": qty,
            "unit": item["unit"],
            "unit_price": unit_price,
            "subtotal": round(qty * unit_price, 2),
        })

    result.sort(key=lambda x: x["material"].name)
    return result


@router.get(
    "/plans/{plan_id}/materials-summary",
    response_model=list[MaterialSummaryItem],
)
def get_materials_summary(
    plan_id: int,
    page: int | None = Query(None, ge=1),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> list[dict]:
    plan = db.get(Plan, plan_id)
    if plan is None or plan.project.organization_id != user.organization_id:
        raise HTTPException(status_code=404, detail="Plan no encontrado")
    return _get_materials_summary_data(plan_id, page, db)


@router.get("/plans/{plan_id}/export/xlsx")
def export_xlsx(
    plan_id: int,
    page: int | None = Query(None, ge=1),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> StreamingResponse:
    plan = db.get(Plan, plan_id)
    if plan is None or plan.project.organization_id != user.organization_id:
        raise HTTPException(status_code=404, detail="Plan no encontrado")

    summary_items = _get_materials_summary_data(plan_id, page, db)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Presupuesto"
    ws.views.sheetView[0].showGridLines = True

    title_font = Font(name="Segoe UI", size=16, bold=True, color="1B365D")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=11, color="000000")
    total_font = Font(name="Segoe UI", size=11, bold=True, color="1B365D")
    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    total_fill = PatternFill(start_color="E9EFF5", end_color="E9EFF5", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )
    total_border = Border(
        top=Side(style="thin", color="1B365D"),
        bottom=Side(style="double", color="1B365D"),
    )

    plan_label = plan.project.name if plan.project else plan.original_filename
    ws.merge_cells("A1:E1")
    ws["A1"] = f"Presupuesto de Materiales - {plan_label}"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 40

    headers = ["Material", "Cantidad", "Unidad", "Precio Unitario", "Subtotal"]
    ws.row_dimensions[3].height = 25
    for col_idx, text in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center" if col_idx > 1 else "left", vertical="center"
        )
        cell.border = thin_border

    current_row = 4
    total_budget = 0.0
    for item in summary_items:
        ws.row_dimensions[current_row].height = 20
        cells = [
            (item["material"].name, "left"),
            (item["quantity"], "right"),
            (item["unit"], "center"),
            (item["unit_price"], "right"),
            (item["subtotal"], "right"),
        ]
        for col_idx, (value, align) in enumerate(cells, 1):
            c = ws.cell(row=current_row, column=col_idx, value=value)
            c.font = data_font
            c.alignment = Alignment(horizontal=align, vertical="center")
            c.border = thin_border
            if col_idx in (2, 4, 5):
                c.number_format = "$#,##0.00" if col_idx in (4, 5) else "#,##0.00"
        total_budget += item["subtotal"]
        current_row += 1

    ws.row_dimensions[current_row].height = 25
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    for col_idx in range(1, 5):
        c = ws.cell(row=current_row, column=col_idx)
        c.border = total_border
        c.fill = total_fill
    total_label = ws.cell(row=current_row, column=1, value="Total General")
    total_label.font = total_font
    total_label.alignment = Alignment(horizontal="right", vertical="center")
    total_val = ws.cell(row=current_row, column=5, value=total_budget)
    total_val.font = total_font
    total_val.number_format = "$#,##0.00"
    total_val.alignment = Alignment(horizontal="right", vertical="center")
    total_val.border = total_border
    total_val.fill = total_fill

    for col in ws.columns:
        max_len = max(
            (len(str(cell.value)) for cell in col if cell.row != 1 and cell.value),
            default=0,
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 4, 12)
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width, 30)

    out = BytesIO()
    wb.save(out)
    out.seek(0)

    filename = f"presupuesto_{plan_label.replace(' ', '_')}.xlsx"
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/plans/{plan_id}/export/pdf")
def export_pdf(
    plan_id: int,
    page: int | None = Query(None, ge=1),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> StreamingResponse:
    plan = db.get(Plan, plan_id)
    if plan is None or plan.project.organization_id != user.organization_id:
        raise HTTPException(status_code=404, detail="Plan no encontrado")

    summary_items = _get_materials_summary_data(plan_id, page, db)

    from datetime import datetime

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40
    )
    styles = getSampleStyleSheet()

    def _style(name, parent="Normal", **kw):
        return ParagraphStyle(name, parent=styles[parent], **kw)

    title_s = _style("T", "Heading1", fontName="Helvetica-Bold", fontSize=20, leading=24,
                     textColor=colors.HexColor("#1B365D"), spaceAfter=15)
    sub_s = _style("S", fontName="Helvetica", fontSize=10, leading=14,
                   textColor=colors.HexColor("#666666"), spaceAfter=25)
    body_s = _style("B", fontName="Helvetica", fontSize=10, leading=12,
                    textColor=colors.HexColor("#333333"))
    body_r = _style("BR", fontName="Helvetica", fontSize=10, leading=12,
                    textColor=colors.HexColor("#333333"), alignment=2)
    body_rb = _style("BRB", fontName="Helvetica-Bold", fontSize=10, leading=12,
                     textColor=colors.HexColor("#333333"), alignment=2)
    hdr_s = _style("H", fontName="Helvetica-Bold", fontSize=10, leading=12,
                   textColor=colors.white)
    hdr_r = _style("HR", fontName="Helvetica-Bold", fontSize=10, leading=12,
                   textColor=colors.white, alignment=2)

    plan_label = plan.project.name if plan.project else plan.original_filename
    date_str = datetime.now().strftime("%d/%m/%Y %H:%M")
    page_txt = f" - Página {page}" if page else ""
    story = [
        Paragraph("Presupuesto de Materiales", title_s),
        Paragraph(f"Plano: {plan_label}{page_txt}<br/>Generado el: {date_str}", sub_s),
    ]

    table_data = [[
        Paragraph("Material", hdr_s),
        Paragraph("Cantidad", hdr_r),
        Paragraph("Unidad", hdr_s),
        Paragraph("Precio Unitario", hdr_r),
        Paragraph("Subtotal", hdr_r),
    ]]

    def _fmt(n: float, prefix: str = "") -> str:
        return f"{prefix}{n:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    total_budget = 0.0
    for item in summary_items:
        table_data.append([
            Paragraph(item["material"].name, body_s),
            Paragraph(_fmt(item["quantity"]), body_r),
            Paragraph(item["unit"], body_s),
            Paragraph(_fmt(item["unit_price"], "$"), body_r),
            Paragraph(_fmt(item["subtotal"], "$"), body_r),
        ])
        total_budget += item["subtotal"]

    table_data.append([
        Paragraph("Total General", body_rb), "", "", "",
        Paragraph(_fmt(total_budget, "$"), body_rb),
    ])

    t = Table(table_data, colWidths=[240, 70, 50, 80, 92])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
        ("GRID", (0, 0), (-1, -2), 0.5, colors.HexColor("#E2E8F0")),
        ("SPAN", (0, -1), (3, -1)),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F8FAFC")),
        ("TOPPADDING", (0, -1), (-1, -1), 10),
        ("BOTTOMPADDING", (0, -1), (-1, -1), 10),
        ("LINEABOVE", (0, -1), (-1, -1), 1.5, colors.HexColor("#1B365D")),
        ("LINEBELOW", (0, -1), (-1, -1), 2, colors.HexColor("#1B365D")),
    ]))
    story.append(t)
    doc.build(story)
    buffer.seek(0)

    filename = f"presupuesto_{plan_label.replace(' ', '_')}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
